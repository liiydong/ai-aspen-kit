# -*- coding: utf-8 -*-
import openpyxl, sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

SRC = r'C:\Users\<用户名>\Desktop\ASPEN PLUS-最终版报告-2026.xlsx'
OUT = r'D:\<化工工作区>\_probe\report_dump.txt'

wb = openpyxl.load_workbook(SRC, data_only=True)
lines = []
lines.append('SHEETS: ' + ' | '.join(wb.sheetnames))
for ws in wb.worksheets:
    lines.append('')
    lines.append('='*100)
    lines.append('SHEET: %s   dims=%s  max_row=%d max_col=%d' % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    lines.append('='*100)
    for r in range(1, ws.max_row+1):
        cells = []
        for c in range(1, ws.max_column+1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            s = str(v).strip()
            if s == '':
                continue
            cells.append(s)
        if cells:
            lines.append('R%04d| %s' % (r, ' || '.join(cells)))

with open(OUT, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print('done rows written:', len(lines))
