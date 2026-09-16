# -*- coding: utf-8 -*-
"""
verify_report_parse.py —— 案例复现自检脚本

用途：验证「用本仓库记载的方法解析 Aspen 计算报告」能否得到
      《数据核对与订正说明.md》里记录的订正值。

用法（需要 openpyxl）：
    python verify_report_parse.py

无需 Aspen 授权 —— 只读 xlsx，不碰 .bkp。

本脚本刻意使用与 04-脚本库/dump_cols.py + parse_report.py 完全相同的
算法（A/B/C 三列重建、三个解析陷阱的处理），以此证明那条复现路径真实可用。
"""
import re
import os
import sys
import io
import collections

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    import openpyxl
except ImportError:
    sys.exit('需要 openpyxl：pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'ASPEN PLUS-最终版报告-2026.xlsx')

if not os.path.exists(XLSX):
    sys.exit('找不到报告文件：%s' % XLSX)

# ---------- 第 1 步：xlsx → A/B/C 三列（同 dump_cols.py）----------
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Sheet1']
lines = []
for r in range(1, ws.max_row + 1):
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        vals.append('' if v is None else str(v))
    if all(x.strip() == '' for x in vals):
        continue
    lines.append('%d\t%s' % (r, '\t'.join(vals)))

print('=' * 66)
print('案例复现自检')
print('=' * 66)
print('报告文件      : %s' % os.path.basename(XLSX))
print('非空行数      : %d' % len(lines))
print('列数          : %d  ← 应为 3（Aspen 按固定字符位切的 A/B/C 三列）' % ws.max_column)

# ---------- 第 2 步：重建固定宽度行（同 parse_report.py）----------
COMPONENTS = ['3-MP', 'NH3', 'O2', 'N2', '3-CP', '4-CP', 'H2O', 'TOL',
              'CO2', 'HCN', 'CO', 'AIR', 'NAM', 'NAC', '4-MP']
NUM = re.compile(r'[-+]?(?:\d+\.\d+|\d+)(?:[-+]\d+)?')
STREAMID = re.compile(r'S-\d+[A-Z]?')
PAGEHDR = re.compile(r'^S-\d+[A-Z]?(?:\s+S-\d+[A-Z]?)*$')


def nums(s, n):
    toks = NUM.findall(s)
    if len(toks) < n:
        return None
    return [float_asp(x) for x in toks[-n:]]


def float_asp(t):
    """Aspen 数值写法：1.2345-02 → 0.012345"""
    m = re.match(r'^([-+]?)(\d+\.\d+|\d+)(?:([-+])(\d+))?$', t)
    if not m:
        return None
    sign, mant, esign, exp = m.groups()
    v = float(mant)
    if esign:
        v = v * (10 ** (-int(exp) if esign == '-' else int(exp)))
    return -v if sign == '-' else v


raw = []
for ln in lines:
    p = ln.split('\t')
    r = int(p[0])
    cols = (p[1:] + ['', '', ''])[:3]
    a, b, c = cols[0], cols[1], cols[2]
    # ★陷阱 2★ B→C 切点会吃掉一个空格；不补会让两个数值粘连（13.9059 + 13.9059 → 13.905913.9059）
    if b and c and not b.endswith(('-', '+')):
        b = b + ' '
    raw.append((r, (a + b + c).rstrip()))

# ---------- 第 3 步：解析流股 ----------
streams = collections.OrderedDict()
cur_ids = []
for r, s in raw:
    st = s.strip()
    # ★陷阱 4★ 页首行「S-101 S-102 …」才是可靠的分页标记
    if PAGEHDR.match(st) and 'S-' in st and len(st) < 60:
        ids = STREAMID.findall(st)
        if len(ids) >= 1:
            cur_ids = ids
            for sid in ids:
                streams.setdefault(sid, {})
        continue
    if st.startswith('BLOCK:'):
        cur_ids = []
        continue
    n = len(cur_ids)
    if n == 0:
        continue
    # ★陷阱 3★ 组分名里的数字会污染取值（H2O 的 2、3-CP 的 3）→ 先剥离标签
    for comp in COMPONENTS:
        if st.startswith(comp) and (len(st) == len(comp) or not st[len(comp)].isalnum()):
            seg = st[len(comp):]
            v = nums(seg, n) or nums(s, n)
            if v:
                for i, sid in enumerate(cur_ids):
                    streams[sid].setdefault('comp_kmol', {})[comp] = v[i]
            break
    else:
        key = None
        for tag, name in [
            ('TEMP', 'TEMP'), ('PRES', 'PRES'), ('VFRAC', 'VFRAC'),
            ('LFRAC', 'LFRAC'), ('SFRAC', 'SFRAC'),
            ('KCAL/MOL', 'H_KCALMOL'), ('KCAL/KG', 'H_KCALKG'), ('GCAL/HR', 'H_GCALHR'),
            ('CAL/MOL-K', 'S_CALMOLK'), ('CAL/GM-K', 'S_CALGMK'),
            ('KMOL/CUM', 'RHO_KMOLCUM'), ('KG/CUM', 'RHO_KGCUM'),
            ('AVG MW', 'MW'),
            ('KMOL/HR', 'F_KMOL'), ('KG/HR', 'F_KG'), ('CUM/HR', 'F_CUM'),
        ]:
            if st.startswith(tag) and (len(st) == len(tag) or not st[len(tag)].isalnum()):
                key = name
                break
        if key:
            v = nums(s, n)
            if v:
                for i, sid in enumerate(cur_ids):
                    streams[sid][key] = v[i]

print('解析出流股数  : %d  ← 文档记载 61 条' % len(streams))
print('带流量数据的  : %d' % sum(1 for d in streams.values() if 'F_KG' in d))

# ---------- 第 4 步：比对《数据核对与订正说明.md》的订正值 ----------
EXPECT_KG = collections.OrderedDict([
    ('S-209', ('产品（烟酰胺）', 1411.6)),
    ('S-110', ('吸收液', 3464.6)),
    ('S-113', ('萃余水', 2172.4)),
    ('S-114', ('萃取相', 3699.3)),
    ('S-115', ('回收甲苯', 2333.6)),
    ('S-116', ('T-401 塔底', 1365.7)),
    ('S-118', ('T-402 塔底', 1273.8)),
    ('S-121', ('3-CP 产品', 1254.5)),
])

print('\n' + '-' * 66)
print('与《数据核对与订正说明.md》记录的订正值比对（单位 kg/h）')
print('-' * 66)
print('%-8s %-16s %12s %12s %6s' % ('流股', '名称', '文档值', '解析值', '判定'))

passed = failed = 0
for sid, (name, expect) in EXPECT_KG.items():
    got = streams.get(sid, {}).get('F_KG')
    if got is None:
        print('%-8s %-16s %12.1f %12s %6s' % (sid, name, expect, '未取到', '⚠'))
        failed += 1
        continue
    ok = abs(got - expect) / expect < 0.001   # 容差 0.1%（报告打印位数舍入）
    print('%-8s %-16s %12.1f %12.1f %6s' % (sid, name, expect, got, '✅' if ok else '❌'))
    passed += 1 if ok else 0
    failed += 0 if ok else 1

# ---------- 汇总 ----------
print('\n' + '=' * 66)
print('结果：%d 项对上，%d 项不符' % (passed, failed))
print('=' * 66)
if failed == 0 and len(streams) >= 50:
    print('\n✅ 复现成功：按本仓库记载的方法解析，得到的数字与文档一致。')
    print('   说明这条复现路径真实可用。')
    sys.exit(0)
print('\n⚠ 有出入，请对照《数据核对与订正说明.md》人工核对。')
sys.exit(1)
